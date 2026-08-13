#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""解析 OA 立项附件：xls/xlsx（pandas）+ docx（python-docx），输出关键内容。
用法: python scripts/parse_attachments.py [目录]
"""
from __future__ import annotations

import sys
from pathlib import Path


def read_excel(path: Path, max_rows: int = 40) -> str:
    import pandas as pd
    try:
        df = pd.read_excel(path, header=None)
    except Exception:
        df = pd.read_excel(path)
    lines = []
    for i, row in enumerate(df.head(max_rows).itertuples(index=False)):
        cells = [str(c)[:50] for c in row if not pd.isna(c)]
        if cells:
            lines.append(f"  R{i}: " + " | ".join(cells))
    return "\n".join(lines)


def read_docx(path: Path, max_paras: int = 90, max_tab_rows: int = 25) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = []
    para_count = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t[:120])
            para_count += 1
            if para_count >= max_paras:
                break
    for ti, table in enumerate(doc.tables[:3]):
        parts.append(f"[表格{ti + 1}]")
        for r, row in enumerate(table.rows[:max_tab_rows]):
            cells = [c.text.strip().replace("\n", " ")[:40] for c in row.cells]
            parts.append("  " + " | ".join(cells))
    return "\n".join(parts)


def main(directory: str):
    d = Path(directory)
    files = sorted(d.glob("*"))
    for f in files:
        if not f.is_file():
            continue
        print(f"\n{'=' * 70}\n### {f.name}（{f.stat().st_size} 字节）")
        try:
            if f.suffix.lower() in (".xls", ".xlsx"):
                print(read_excel(f))
            elif f.suffix.lower() == ".docx":
                print(read_docx(f))
            else:
                print("  (跳过：非 excel/word)")
        except Exception as e:
            print(f"  解析失败: {e}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(Path.home() / ".easybrowser/downloads/attachments"))
